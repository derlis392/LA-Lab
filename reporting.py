"""LA-Lab: reportes reproducibles y narrativas cautelosas para educación."""
REPORT_VERSION = "0.6.11"
import io
from datetime import datetime
from xml.sax.saxutils import escape
import numpy as np
import pandas as pd


def describe_data(df, question='', target=None):
    n,m=df.shape; missing=int(df.isna().sum().sum()); dup=int(df.duplicated().sum())
    nums=df.select_dtypes(include='number').columns
    txt=f'Se analizaron {n:,} observaciones y {m} variables ({len(nums)} numéricas). Se detectaron {missing} valores faltantes y {dup} filas duplicadas.'
    if question: txt+=f' Pregunta orientadora: {question}'
    if target in df:
        s=df[target]
        if pd.api.types.is_numeric_dtype(s):
            txt+=f' La variable objetivo {target} presenta media {s.mean():.2f}, mediana {s.median():.2f} y rango {s.min():.2f}–{s.max():.2f}.'
        else:
            txt+=f' La variable objetivo {target} tiene {s.nunique()} categorías.'
    return txt+' Estos indicadores describen el conjunto; no demuestran aprendizaje ni causalidad.'


def explain_ml(ml):
    if not ml: return 'No se ejecutó un experimento de minería de datos en esta sesión.'
    task=ml['tarea']; met=ml.get('metricas',{}); comp=ml.get('comparacion')
    if task=='Regresión':
        mae=met['MAE']; base=met['Baseline MAE']; r2=met['R2']
        relation='menor' if mae<base else 'mayor o igual'
        improvement=(base-mae)/base*100 if base else 0
        text=(f'El modelo de regresión {ml["algoritmo"]} obtuvo MAE={mae:.3f}, RMSE={met["RMSE"]:.3f} y R²={r2:.3f}. '
              f'El error absoluto fue {relation} que el del modelo de referencia (MAE={base:.3f}); cambio relativo: {improvement:+.1f}% de reducción del error. '
              'La gráfica real-predicho permite identificar sobreestimaciones y subestimaciones; los residuos muestran casos con errores elevados. ')
        if r2<.3: text+='El R² es bajo en esta partición: la capacidad predictiva es limitada. '
        text+='Sugerencia: revisar variables, su disponibilidad temporal y estabilidad mediante validación cruzada. No utilizar para calificar estudiantes automáticamente.'
    elif task=='Clasificación':
        acc=met['Accuracy'];base=met['Baseline accuracy'];bacc=met['Balanced accuracy'];f1=met['F1 macro']
        text=(f'El clasificador {ml["algoritmo"]} alcanzó accuracy={acc:.3f}, balanced accuracy={bacc:.3f} y F1 macro={f1:.3f}. '
              f'El clasificador mayoritario de referencia obtuvo accuracy={base:.3f}. ')
        text+=('La accuracy es inferior a la referencia; no debe interpretarse como una mejora global. ' if acc<=base else 'La accuracy supera la referencia, pero debe examinarse el desempeño por clase. ')
        classes=ml.get('por_clase',[])
        if classes:
            minority=min(classes,key=lambda row:row['support'])
            text+=f'La clase con menor soporte en prueba ({minority["Clase"]}, n={int(minority["support"])}) obtuvo precisión={minority["precision"]:.3f}, recall={minority["recall"]:.3f} y F1={minority["f1-score"]:.3f}. '
            if minority['support']<20:text+='La estimación de esa clase es inestable por su reducido número de casos. '
        text+='Sugerencia: revisar errores por clase, desbalance y criterios de selección; no utilizar etiquetas predictivas para decisiones individuales.'
    else:
        sil=met.get('Silhouette');counts=ml.get('conteos',{})
        text=f'El agrupamiento {ml["algoritmo"]} produjo {len(counts)} etiquetas con tamaños '+', '.join(f'{k}: {v}' for k,v in counts.items())+'. '
        if sil is not None:
            text+=f'Silhouette={sil:.3f}. '
            if sil<.25:text+='La separación observada es débil. '
            elif sil<.5:text+='La separación es modesta y requiere contrastes adicionales. '
        if ml.get('comparacion_k'):
            best=max(ml['comparacion_k'],key=lambda row:row['silhouette'])
            text+=f'En la exploración de k, el valor más alto corresponde a k={best["k"]} (silhouette={best["silhouette"]:.3f}). '
        if ml.get('perfiles'):
            profiles=ml['perfiles']; keys=[k for k in profiles[0] if k!='Grupo']
            for row in profiles:
                if keys:
                    text+=f'Grupo {row["Grupo"]}: '+', '.join(f'{k}={row[k]:.2f}' for k in keys[:5])+'. '
            text+='Estos promedios describen variables codificadas; no son diagnósticos. '
        text+='La proyección PCA es una representación aproximada; los grupos no constituyen perfiles educativos validados. Sugerencia: caracterizar variables por grupo, examinar estabilidad y contrastar con teoría pedagógica.'
    if comp:
        metric='F1 macro' if task=='Clasificación' else 'MAE'
        best=(max(comp,key=lambda r:r[metric]) if task=='Clasificación' else min(comp,key=lambda r:r[metric]))
        text+=f' En la comparación exploratoria, {best["Algoritmo"]} obtuvo el mejor {metric} ({best[metric]:.3f}) en esta partición; seleccionar formalmente con validación en entrenamiento, no optimizando sobre prueba.'
    cv=ml.get('validacion_cruzada')
    if cv:
        vals=cv['valores'];text+=f' Validación cruzada en entrenamiento: {cv["metrica"]}={np.mean(vals):.3f} ± {np.std(vals):.3f} ({len(vals)} particiones).'
    return text


def summarize_plot(df, plot):
    if not plot:return 'No se registró una visualización exploratoria específica; las figuras descriptivas del informe se generan con los datos actuales.'
    kind=plot['tipo']
    if kind=='dispersion':
        x,y=plot['x'],plot['y'];pair=df[[x,y]].dropna();r=pair[x].corr(pair[y])
        return f'La correlación lineal de Pearson entre {x} y {y} es r={r:.3f} (n={len(pair)}). Una asociación no prueba causalidad ni descarta confusores.'
    if kind=='correlacion':return 'La matriz de correlación permite identificar asociaciones lineales; se deben comprobar escalas, codificaciones y relaciones mecánicas entre variables.'
    if kind=='histograma':return f'La distribución de {plot["x"]} describe frecuencia, dispersión y posibles extremos; su interpretación requiere conocer la escala.'
    if kind=='caja':return f'El diagrama de caja de {plot["y"]} permite observar mediana, dispersión y posibles extremos, sin inferir efectos causales.'
    return 'El gráfico describe frecuencias por categoría, no relaciones causales.'


def _chart_image(df,plot,ml,which):
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    fig,ax=plt.subplots(figsize=(8,4.3),dpi=120)
    if which=='correlation':
        cols=list(df.select_dtypes(include='number').columns[:10])
        target=ml.get('objetivo') if ml else None
        if target in df.select_dtypes(include='number') and target not in cols: cols=cols[:9]+[target]
        if len(cols)<2:plt.close(fig);return None
        corr=df[cols].corr();im=ax.imshow(corr,vmin=-1,vmax=1,cmap='coolwarm')
        ax.set_xticks(range(len(cols)),cols,rotation=65,ha='right',fontsize=7)
        ax.set_yticks(range(len(cols)),cols,fontsize=7);fig.colorbar(im,ax=ax,fraction=.04,pad=.02)
        ax.set_title('Matriz de correlación (Pearson)')
    elif which=='overview':
        numeric=df.select_dtypes(include='number').columns
        target=ml.get('objetivo') if ml else None
        if len(numeric):
            col=target if target in numeric else numeric[0]; ax.hist(df[col].dropna(),bins=min(20,max(5,df[col].nunique())),edgecolor='white');ax.set_title('Distribución de '+str(col));ax.set_xlabel(str(col));ax.set_ylabel('Frecuencia')
        else:
            col=df.columns[0];v=df[col].astype(str).value_counts().head(10);ax.bar(v.index,v.values);ax.tick_params(axis='x',rotation=45);ax.set_title('Frecuencias de '+str(col))
    elif which=='plot' and plot:
        kind=plot['tipo']
        if kind=='dispersion':ax.scatter(df[plot['x']],df[plot['y']],alpha=.5,s=14);ax.set_xlabel(plot['x']);ax.set_ylabel(plot['y'])
        elif kind=='histograma':ax.hist(df[plot['x']].dropna(),bins=20);ax.set_xlabel(plot['x'])
        elif kind=='correlacion':plt.close(fig);return _chart_image(df,plot,ml,'correlation')
        elif kind=='caja':
            y=plot['y'];group=plot['grupo']
            if group!='Sin grupos' and group in df:
                names=df[group].dropna().astype(str).value_counts().head(8).index
                ax.boxplot([df.loc[df[group].astype(str)==name,y].dropna() for name in names],tick_labels=names);ax.tick_params(axis='x',rotation=35)
            else:ax.boxplot(df[y].dropna());ax.set_ylabel(y)
        elif kind=='categorias':
            v=df[plot['x']].astype(str).value_counts().head(12);ax.bar(v.index,v.values);ax.tick_params(axis='x',rotation=45)
        ax.set_title('Visualización exploratoria registrada')
    elif which=='ml_primary' and ml:
        task=ml['tarea']
        if task=='Regresión':
            real=np.array(ml['reales'],dtype=float);pred=np.array(ml['predicciones'],dtype=float)
            ax.scatter(real,pred,alpha=.65,s=16);low=min(real.min(),pred.min());high=max(real.max(),pred.max());ax.plot([low,high],[low,high],'--',label='Predicción ideal');ax.set(xlabel='Real',ylabel='Predicho',title='Valores reales y predichos');ax.legend()
        elif task=='Clasificación':
            cm=np.array(ml['matriz']);ax.imshow(cm,cmap='Blues');names=ml['etiquetas'];ax.set_xticks(range(len(names)),names,rotation=45);ax.set_yticks(range(len(names)),names)
            for i in range(len(names)):
                for j in range(len(names)):ax.text(j,i,str(cm[i,j]),ha='center',va='center',color='white' if cm[i,j]>cm.max()/2 else 'black')
            ax.set(xlabel='Predicción',ylabel='Clase real',title='Matriz de confusión')
        else:
            coords=np.array(ml['pca_coordenadas']);labels=np.array(ml['etiquetas_grupo']);
            for label in np.unique(labels):
                mask=labels==label;ax.scatter(coords[mask,0],coords[mask,1],s=13,alpha=.7,label=str(label))
            ax.legend(title='Grupo',fontsize=8);ax.set(xlabel='Componente 1',ylabel='Componente 2',title='Agrupamiento proyectado en PCA')
    elif which=='ml_secondary' and ml:
        task=ml['tarea']
        if task=='Regresión':
            real=np.array(ml['reales'],dtype=float);pred=np.array(ml['predicciones'],dtype=float)
            ax.scatter(pred,real-pred,s=16,alpha=.65);ax.axhline(0,linestyle='--',color='gray');ax.set(xlabel='Predicción',ylabel='Residuo (real - predicho)',title='Diagnóstico de residuos')
        elif task=='Clasificación':
            records=ml.get('por_clase',[]);x=np.arange(len(records));w=.26
            for j,metric in enumerate(['precision','recall','f1-score']):ax.bar(x+(j-1)*w,[r[metric] for r in records],w,label=metric)
            ax.set_xticks(x,[r['Clase'] for r in records]);ax.set_ylim(0,1);ax.legend();ax.set_title('Métricas por clase')
        else:
            vals=ml.get('comparacion_k')
            if vals:ax.plot([r['k'] for r in vals],[r['silhouette'] for r in vals],marker='o');ax.set(xlabel='k',ylabel='Silhouette',title='Comparación de número de grupos')
            else:
                counts=ml.get('conteos',{});ax.bar(list(map(str,counts.keys())),list(counts.values()));ax.set(xlabel='Grupo',ylabel='Observaciones',title='Tamaños de grupo')
    elif which=='comparison' and ml and ml.get('comparacion'):
        rows=ml['comparacion'];metric='F1 macro' if ml['tarea']=='Clasificación' else 'MAE'
        ax.barh([r['Algoritmo'] for r in rows],[r[metric] for r in rows]);ax.invert_yaxis();ax.set(xlabel=metric,title='Comparación exploratoria de algoritmos')
    else:plt.close(fig);return None
    fig.tight_layout();buf=io.BytesIO();fig.savefig(buf,format='png',dpi=125,bbox_inches='tight');plt.close(fig);buf.seek(0);return buf


# Diccionario interpretativo limitado a variables documentadas de Student Performance.
STUDENT_VARIABLES = {
 'school':('Escuela de procedencia','Nominal','GP/MS'),
 'sex':('Sexo registrado','Nominal','F/M'),
 'age':('Edad','Cuantitativa','Años'),
 'address':('Tipo de residencia','Nominal','U/R'),
 'famsize':('Tamaño del hogar','Nominal','LE3/GT3'),
 'Pstatus':('Situación de convivencia parental','Nominal','T/A'),
 'Medu':('Escolaridad materna','Ordinal','0–4'),
 'Fedu':('Escolaridad paterna','Ordinal','0–4'),
 'Mjob':('Ocupación materna','Nominal','Categorías UCI'),
 'Fjob':('Ocupación paterna','Nominal','Categorías UCI'),
 'reason':('Motivo de elección de escuela','Nominal','Categorías UCI'),
 'guardian':('Tutor responsable','Nominal','Categorías UCI'),
 'traveltime':('Tiempo de traslado','Ordinal','1–4'),
 'studytime':('Tiempo semanal de estudio','Ordinal','1–4'),
 'failures':('Reprobaciones previas','Conteo codificado','0–3; 3 puede indicar 3 o más'),
 'schoolsup':('Apoyo educativo escolar','Binaria','yes/no'),
 'famsup':('Apoyo educativo familiar','Binaria','yes/no'),
 'paid':('Clases adicionales pagadas','Binaria','yes/no'),
 'activities':('Actividades extracurriculares','Binaria','yes/no'),
 'nursery':('Asistió a educación preescolar','Binaria','yes/no'),
 'higher':('Desea educación superior','Binaria','yes/no'),
 'internet':('Acceso a Internet en casa','Binaria','yes/no'),
 'romantic':('Relación romántica registrada','Binaria','yes/no'),
 'famrel':('Calidad de relaciones familiares','Ordinal','1–5'),
 'freetime':('Tiempo libre después de escuela','Ordinal','1–5'),
 'goout':('Frecuencia de salidas','Ordinal','1–5'),
 'Dalc':('Consumo de alcohol entre semana','Ordinal','1–5'),
 'Walc':('Consumo de alcohol en fin de semana','Ordinal','1–5'),
 'health':('Estado de salud autopercibido','Ordinal','1–5'),
 'absences':('Número de ausencias escolares','Conteo','Días registrados'),
 'G1':('Calificación del primer periodo','Cuantitativa','0–20'),
 'G2':('Calificación del segundo periodo','Cuantitativa','0–20'),
 'G3':('Calificación final','Cuantitativa','0–20')
}
def variable_dictionary(df, meta):
    is_student='Student Performance' in str(meta.get('fuente',''))
    return [{'Variable':c,'Significado':STUDENT_VARIABLES[c][0] if is_student and c in STUDENT_VARIABLES else 'Consultar diccionario original',
             'Tipo conceptual':STUDENT_VARIABLES[c][1] if is_student and c in STUDENT_VARIABLES else ('Numérica sin escala validada' if pd.api.types.is_numeric_dtype(df[c]) else 'Categoría sin escala validada'),
             'Escala / valores':STUDENT_VARIABLES[c][2] if is_student and c in STUDENT_VARIABLES else 'Ver documentación de la fuente'} for c in df.columns]

def research_note(meta, ml):
    if not ml:return 'No se ejecutó minería de datos; la pregunta orientadora permanece como guía exploratoria.'
    vars=set(ml.get('variables',[])); target=ml.get('objetivo')
    if 'Student Performance' in str(meta.get('fuente','')) and target=='G3':
        if vars.intersection({'G1','G2'}):return 'El modelo incluye calificaciones de periodos anteriores: sus resultados corresponden a un escenario posterior a esas evaluaciones, no a una predicción inicial del curso.'
        return 'La pregunta orientadora se refiere a evaluaciones y calificación final, pero el modelo no incorpora G1 ni G2. El experimento solo estima la relación predictiva de los predictores seleccionados con G3.'
    return 'La pregunta orientadora y el experimento configurado deben contrastarse antes de derivar conclusiones educativas.'

def build_pdf(df,meta,ml=None,plot=None):
    from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,Image,Table,TableStyle,KeepTogether,CondPageBreak,KeepInFrame
    from reportlab.lib.styles import getSampleStyleSheet,ParagraphStyle
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.units import cm
    from reportlab.lib.pagesizes import A4
    buf=io.BytesIO();doc=SimpleDocTemplate(buf,pagesize=A4,leftMargin=1.8*cm,rightMargin=1.8*cm,topMargin=1.7*cm,bottomMargin=1.8*cm)
    styles=getSampleStyleSheet();styles.add(ParagraphStyle(name='LabTitle',parent=styles['Title'],fontSize=19,leading=23,textColor=colors.HexColor('#15375b'),spaceAfter=14))
    styles.add(ParagraphStyle(name='LabBody',parent=styles['BodyText'],fontSize=9.5,leading=14,spaceAfter=8))
    styles.add(ParagraphStyle(name='LabH',parent=styles['Heading2'],textColor=colors.HexColor('#15375b'),spaceBefore=13,spaceAfter=7,keepWithNext=True))
    story=[]
    def para(text):story.append(Paragraph(escape(str(text)).replace('\n','<br/>'),styles['LabBody']))
    def heading(text):story.append(CondPageBreak(2.3*cm));story.append(Paragraph(escape(text),styles['LabH']))
    def table(rows):
        cells=[[Paragraph(escape(str(x)),styles['LabBody']) for x in row] for row in rows]
        t=Table(cells,repeatRows=1,hAlign='LEFT',colWidths=None,splitByRow=1)
        t.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#e7eef6')),('GRID',(0,0),(-1,-1),.25,colors.HexColor('#cbd5e1')),('VALIGN',(0,0),(-1,-1),'TOP'),('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,colors.HexColor('#f7f9fc')]),('LEFTPADDING',(0,0),(-1,-1),6),('RIGHTPADDING',(0,0),(-1,-1),6)]))
        story.append(t);story.append(Spacer(1,9))
    def figure(which,caption):
        png=_chart_image(df,plot,ml,which)
        if png:
            from PIL import Image as PILImage
            im=PILImage.open(png);w,h=im.size; maxw=16.5*cm;maxh=6.8*cm;scale=min(maxw/w,maxh/h)
            png.seek(0)
            story.append(KeepTogether([Image(png,width=w*scale,height=h*scale),Paragraph(escape(caption),styles['LabBody'])]))
    from pathlib import Path
    assets=Path(__file__).resolve().parent/'assets'
    logo_cells=[]
    logo_specs=[('unam.png',2.0*cm,2.0*cm),('cuaed.png',5.0*cm,1.8*cm),('linea_investigacion.svg',2.3*cm,2.0*cm)]
    for file,max_width,max_height in logo_specs:
        path=assets/file
        if not path.exists():continue
        from PIL import Image as PILImage
        if path.suffix.lower()=='.svg':
            # Rasterización a alta resolución únicamente para incrustar en PDF.
            # Se conserva el SVG original en assets y en la interfaz web.
            try:
                import cairosvg
                rendered=io.BytesIO(cairosvg.svg2png(url=str(path),output_width=900))
                iw,ih=PILImage.open(rendered).size;rendered.seek(0)
                source=rendered
            except Exception:
                continue
        else:
            iw,ih=PILImage.open(path).size;source=str(path)
        factor=min(max_width/iw,max_height/ih)
        logo_cells.append(Image(source,width=iw*factor,height=ih*factor))
    if logo_cells:
        logo_table=Table([logo_cells],hAlign='CENTER',colWidths=[3.2*cm,6.5*cm,3.2*cm][:len(logo_cells)])
        logo_table.setStyle(TableStyle([('VALIGN',(0,0),(-1,-1),'MIDDLE'),('ALIGN',(0,0),(-1,-1),'CENTER'),('LEFTPADDING',(0,0),(-1,-1),5),('RIGHTPADDING',(0,0),(-1,-1),5)]))
        story.append(logo_table);story.append(Spacer(1,12))
    story.append(Paragraph('LA-Lab v0.6.11 | Informe académico de analítica del aprendizaje',styles['LabTitle']))
    para('Caso: '+meta['caso']);para('Generado: '+datetime.now().strftime('%Y-%m-%d %H:%M')+' | Documento exploratorio; revisión humana requerida')
    heading('1. Propósito y procedencia')
    para('Pregunta de investigación del caso: '+meta['pregunta']);para('Fuente: '+meta['fuente']);para('Licencia declarada: '+meta['licencia'])
    if ml and ml.get('objetivo')=='G3' and 'evaluacion' in __import__('unicodedata').normalize('NFKD',meta['pregunta'].lower()).encode('ascii','ignore').decode('ascii') and not {'G1','G2'}.intersection(ml.get('variables',[])):
        para('Nota de alcance: la pregunta general menciona evaluaciones, pero este experimento no incluye G1 ni G2. Sus resultados describen la asociación predictiva de las variables seleccionadas con G3; no responden directamente a la relación entre evaluaciones.')
    heading('2. Comprensión y calidad de datos');para(describe_data(df,meta['pregunta'],ml.get('objetivo') if ml else None));para('Alcance del experimento: '+research_note(meta,ml))
    para('Transformaciones manuales registradas: '+str(meta['limpieza']).rstrip('.'))
    numeric=df.select_dtypes(include='number');
    if not numeric.empty:
        stat=numeric.describe().T[['mean','std','min','50%','max']].head(12)
        table([['Variable','Media','DE','Mín.','Mediana','Máx.']]+[[str(k)]+[f'{v:.2f}' for v in row] for k,row in stat.iterrows()])
    para('Advertencia sobre escalas: las medias y las correlaciones de Pearson de variables ordinales codificadas (por ejemplo, Medu, Fedu, studytime y health) requieren cautela; los códigos no son intervalos necesariamente equivalentes.')
    heading('2.1 Diccionario de variables (selección)')
    dictionary=variable_dictionary(df,meta)
    important={'age','Medu','Fedu','studytime','failures','absences','G1','G2','G3'}
    selected_dict=[row for row in dictionary if row['Variable'] in important] if 'Student Performance' in str(meta.get('fuente','')) else dictionary[:8]
    if selected_dict:table([['Variable','Significado','Tipo','Escala']]+[[d['Variable'],d['Significado'],d['Tipo conceptual'],d['Escala / valores']] for d in selected_dict])
    heading('3. Visualización e interpretación exploratoria')
    selected = plot.get('x') if plot else None
    overview_var = (ml.get('objetivo') if ml else None) or (str(numeric.columns[0]) if len(numeric.columns) else str(df.columns[0]))
    # Evitar repetir la misma distribución como figura descriptiva y figura seleccionada.
    same_hist = bool(plot and plot.get('tipo')=='histograma' and selected==overview_var)
    if not same_hist:
        figure('overview',f'Figura 1. Distribución de {overview_var}; observar concentración, rango y posibles valores extremos.')
    if len(numeric.columns)>1:
        figure('correlation','Figura 2. Matriz de correlaciones de Pearson entre variables numéricas; las escalas ordinales requieren interpretación cautelosa.')
    if plot:
        if plot.get('tipo')!='correlacion':
            figure('plot','Gráfico exploratorio seleccionado: '+str(plot.get('tipo'))+'.')
        para(summarize_plot(df,plot))
    else:
        para('No se seleccionó un gráfico específico; las figuras descriptivas no sustituyen un análisis contextual.')
    heading('4. Minería de datos y evaluación')
    if ml:
        para('Tarea: '+ml['tarea']+' | Algoritmo: '+ml['algoritmo']+' | Semilla: '+str(ml['semilla'])+' | N: '+str(ml['n']))
        para('Variable objetivo: '+str(ml.get('objetivo'))+' | Predictores: '+', '.join(ml['variables']))
        if ml.get('porcentaje_prueba'):
            para('Partición de prueba: '+str(ml['porcentaje_prueba'])+' %. Entrenamiento: '+str(ml.get('n_entrenamiento','no registrado'))+'; prueba: '+str(ml.get('n_prueba','no registrado'))+'.')
            para('Preprocesamiento interno del modelo: '+str(ml.get('preprocesamiento_interno','Pipeline de imputación/codificación/escalado, ajustado exclusivamente sobre entrenamiento.')))
        metric_rows=[['Métrica','Resultado']]+[[str(k),f'{v:.3f}' if isinstance(v,(int,float)) else str(v)] for k,v in ml.get('metricas',{}).items()]
        story.append(CondPageBreak(min(7.5*cm, (len(metric_rows)+1)*0.75*cm)))
        table(metric_rows)
        if ml['tarea']=='Agrupamiento' and ml.get('perfiles'):
            rows=ml['perfiles']; columns=list(rows[0])[:6]
            table([columns]+[[f'{v:.2f}' if isinstance(v,(int,float)) else str(v) for v in [row[k] for k in columns]] for row in rows])
            para('Las medias caracterizan los grupos en sus escalas originales; las categorías ordinales no representan necesariamente unidades continuas.')
        figure('ml_primary',{'Regresión':'Valores reales frente a predicciones. La línea diagonal indica predicción perfecta.','Clasificación':'Matriz de confusión: filas reales, columnas predichas; examine errores de cada clase.','Agrupamiento':'Proyección PCA de grupos: la separación visual no valida perfiles.'}[ml['tarea']])
        figure('ml_secondary',{'Regresión':'Residuos: los errores grandes o patrones sistemáticos requieren revisión.','Clasificación':'Precisión, recall y F1 por clase: contraste el rendimiento minoritario.','Agrupamiento':'Comparación de silhouette por k o distribución de tamaños de grupos.'}[ml['tarea']])
        if ml.get('comparacion'):
            story.append(CondPageBreak(7.0*cm))
            heading('4.1 Comparación exploratoria de algoritmos')
            para('Las métricas de esta comparación se calcularon sobre el mismo conjunto de prueba únicamente para ilustración. No deben utilizarse para elegir el algoritmo ganador; la selección formal requiere validación cruzada dentro del conjunto de entrenamiento y evaluación final única sobre prueba.')
            rows=ml['comparacion'];keys=list(rows[0]);table([keys]+[[f'{r[k]:.3f}' if isinstance(r[k],(int,float)) else str(r[k]) for k in keys] for r in rows]);figure('comparison','Comparación de algoritmos según la métrica principal del experimento.')
        if ml.get('validacion_cruzada'):
            cv=ml['validacion_cruzada'];para(f'Validación cruzada en entrenamiento ({len(cv["valores"])} particiones): {cv["metrica"]} = {np.mean(cv["valores"]):.3f} ± {np.std(cv["valores"]):.3f}. Valores por partición: '+', '.join(f'{v:.3f}' for v in cv['valores']))
        heading('5. Interpretación contextualizada');para(explain_ml(ml))
        if ml['tarea']=='Regresión':para('MAE expresa el error absoluto promedio en unidades del objetivo; RMSE penaliza más los errores grandes; R² = 1 - SSE/SST, donde SSE es la suma de errores cuadrados y SST es la suma de desviaciones cuadráticas de los valores reales respecto de su media; puede ser negativo. Examine el modelo de referencia y la estabilidad fuera de la muestra.')
        elif ml['tarea']=='Clasificación':para('Accuracy mide la proporción de aciertos; balanced accuracy promedia el recall de las clases; F1 macro promedia los F1 sin ponderar por tamaño. Con clases desbalanceadas, accuracy puede ocultar errores en la minoría.')
        else:para('Silhouette compara cohesión y separación de grupos; valores próximos a cero sugieren superposición. La elección de k exige también estabilidad, interpretabilidad y justificación educativa.')
    else:para('No se registró un experimento de minería de datos en esta sesión. No se reportan métricas ni conclusiones no ejecutadas.')
    heading('6. Recomendaciones y limitaciones')
    para('Interpretación del alcance: '+research_note(meta,ml))
    para('Recomendación docente: formular hipótesis verificables a partir de los patrones observados, contrastar explicaciones alternativas y documentar el contexto, la temporalidad y las escalas de las variables. Utilizar el informe como evidencia de una práctica de análisis, no como diagnóstico automático.')
    para('Limitaciones: el análisis exploratorio no demuestra causalidad; los resultados de una sola partición no acreditan generalización externa. Las variables de rendimiento pueden estar temporalmente relacionadas y producir fuga de información. Los conjuntos educativos requieren controles de privacidad, representatividad y revisión ética.')
    if meta.get('reflexion'):heading('Reflexión de la persona participante');para(meta['reflexion'])
    heading('7. Fuentes, referencias y reproducibilidad')
    para('Conjunto de datos: '+str(meta.get('fuente','No especificado')))
    para('Diccionario de variables: definiciones del caso incluidas en la sección 2.1 y en el libro Excel.')
    if meta.get('fuente_url'):para('Repositorio y descarga original: '+str(meta['fuente_url']))
    para('Licencia declarada: '+str(meta.get('licencia','No especificada; verificar condiciones de uso')))
    if meta.get('articulo'):para('Publicación de referencia: '+str(meta['articulo']))
    if meta.get('articulo_url'):para('Enlace directo a la publicación: '+str(meta['articulo_url']))
    elif meta.get('articulo'):para('Enlace directo a la publicación: pendiente de verificación; no confundir la ficha del conjunto de datos con el artículo.')
    para('Los análisis de LA-Lab son exploratorios y no constituyen una reproducción automática de los resultados publicados.')
    import hashlib, json, platform, importlib.metadata
    fingerprint=hashlib.sha256(pd.util.hash_pandas_object(df,index=True).values.tobytes()).hexdigest()[:16]
    para('Identificador de los datos analizados (SHA-256 abreviado del hash de filas): '+fingerprint+'; filas: '+str(len(df))+'; columnas: '+str(len(df.columns))+'.')
    if ml:
        para('Configuración: tarea='+str(ml.get('tarea'))+'; algoritmo='+str(ml.get('algoritmo'))+'; objetivo='+str(ml.get('objetivo'))+'; predictores='+', '.join(map(str,ml.get('variables',[])))+'; semilla='+str(ml.get('semilla'))+'; porcentaje de prueba='+str(ml.get('porcentaje_prueba','No aplica'))+'.')
        if ml.get('porcentaje_prueba'):
            para('Tamaños efectivos registrados: entrenamiento = '+str(ml.get('n_entrenamiento','no disponible'))+'; prueba = '+str(ml.get('n_prueba','no disponible'))+'.')
        if ml.get('validacion_cruzada'):
            para('Validación cruzada realizada exclusivamente en entrenamiento; consultar métricas y particiones en la sección 4.')
    para('Transformaciones manuales: '+str(meta.get('limpieza','Sin operaciones registradas')).rstrip('.')+'.')
    if ml and ml.get('preprocesamiento_interno'):para('Preprocesamiento supervisado interno: '+str(ml['preprocesamiento_interno']))
    packages=['pandas','numpy','scikit-learn','matplotlib','reportlab','streamlit']
    versions=[]
    for package in packages:
        try:versions.append(package+' '+importlib.metadata.version(package))
        except importlib.metadata.PackageNotFoundError:pass
    para('Entorno: Python '+platform.python_version()+'; '+ '; '.join(versions)+'. Versión de LA-Lab: '+REPORT_VERSION+'.')
    def footer(canvas,doc):
        canvas.setFont('Helvetica',8);canvas.setFillColor(colors.grey);canvas.drawString(1.8*cm,1.1*cm,'LA-Lab v0.6.11 · Informe exploratorio · No tomar decisiones automatizadas sobre estudiantes');canvas.drawRightString(19.2*cm,1.1*cm,str(doc.page))
    doc.build(story,onFirstPage=footer,onLaterPages=footer);return buf.getvalue()


def build_excel(df,meta,ml=None):
    buf=io.BytesIO()
    with pd.ExcelWriter(buf,engine='openpyxl') as writer:
        df.head(5000).to_excel(writer,sheet_name='Datos',index=False)
        pd.DataFrame([{'Campo':k,'Valor':v} for k,v in meta.items()]).to_excel(writer,sheet_name='Metadatos',index=False)
        num=df.select_dtypes(include='number')
        if len(num.columns):
            num.describe().T.to_excel(writer,sheet_name='Estadisticas')
            if len(num.columns)>1:num.corr().to_excel(writer,sheet_name='Correlaciones')
        if ml:
            pd.DataFrame([{'Campo':k,'Valor':str(v)} for k,v in ml.items() if k not in ('reales','predicciones','pca_coordenadas','etiquetas_grupo','por_clase','comparacion','perfiles','comparacion_k')]).to_excel(writer,sheet_name='Experimento',index=False)
            pd.DataFrame([{'Metrica':k,'Valor':v} for k,v in ml.get('metricas',{}).items()]).to_excel(writer,sheet_name='Metricas',index=False)
            if ml.get('comparacion'):pd.DataFrame(ml['comparacion']).to_excel(writer,sheet_name='Comparacion',index=False)
            if ml.get('por_clase'):pd.DataFrame(ml['por_clase']).to_excel(writer,sheet_name='Por_clase',index=False)
            if ml.get('reales') is not None:pd.DataFrame({'Real':ml['reales'],'Prediccion':ml['predicciones']}).to_excel(writer,sheet_name='Predicciones',index=False)
            if ml.get('etiquetas_grupo') is not None:
                pd.DataFrame({'Grupo':ml['etiquetas_grupo'],'PCA_1':[x[0] for x in ml['pca_coordenadas']],'PCA_2':[x[1] for x in ml['pca_coordenadas']]}).to_excel(writer,sheet_name='Grupos',index=False)
            if ml.get('perfiles'):pd.DataFrame(ml['perfiles']).to_excel(writer,sheet_name='Perfiles',index=False)
            if ml.get('comparacion_k'):pd.DataFrame(ml['comparacion_k']).to_excel(writer,sheet_name='Seleccion_k',index=False)
            pd.DataFrame({'Interpretacion':[explain_ml(ml)]}).to_excel(writer,sheet_name='Interpretacion',index=False)
        pd.DataFrame(variable_dictionary(df,meta)).to_excel(writer,sheet_name='Diccionario',index=False)
        import hashlib,platform,importlib.metadata
        fingerprint=hashlib.sha256(pd.util.hash_pandas_object(df,index=True).values.tobytes()).hexdigest()[:16]
        details=[('Versión LA-Lab',REPORT_VERSION),('Fecha de exportación',datetime.now().isoformat(timespec='seconds')),
                 ('Hash abreviado de filas',fingerprint),('Número de filas',len(df)),('Número de columnas',len(df.columns)),
                 ('Versión Python',platform.python_version()),('Pregunta del caso',meta.get('pregunta','')),
                 ('Alcance de resultados',research_note(meta,ml)),('Limitación','Análisis observacional; no implica causalidad ni permite decisiones automatizadas.')]
        for pkg in ['pandas','numpy','scikit-learn','streamlit']:
            try:details.append(('Paquete '+pkg,importlib.metadata.version(pkg)))
            except importlib.metadata.PackageNotFoundError:pass
        if ml:
            for key in ['tarea','algoritmo','objetivo','variables','semilla','porcentaje_prueba','n_entrenamiento','n_prueba','preprocesamiento_interno']:
                details.append(('Modelo: '+key,str(ml.get(key,'No registrado'))))
            if ml.get('validacion_cruzada'):details.append(('Validación cruzada',str(ml['validacion_cruzada'])))
        pd.DataFrame(details,columns=['Campo','Valor']).to_excel(writer,sheet_name='Reproducibilidad',index=False)
        pd.DataFrame([{'Apartado':'Pregunta orientadora','Observación':meta.get('pregunta','')},
          {'Apartado':'Alcance del experimento','Observación':research_note(meta,ml)},
          {'Apartado':'Hallazgos descriptivos','Observación':describe_data(df,meta.get('pregunta',''),ml.get('objetivo') if ml else None)},
          {'Apartado':'Interpretación del modelo','Observación':explain_ml(ml) if ml else 'No se ejecutó modelo'},
          {'Apartado':'Precauciones','Observación':'No inferir causalidad ni equivalencia de intervalos ordinales; validar contexto y temporalidad.'},
          {'Apartado':'Propuesta pedagógica','Observación':'Contrastar hipótesis con docentes y estudiantes; planear una indagación contextualizada antes de cualquier intervención.'}]).to_excel(writer,sheet_name='Sintesis_pedagogica',index=False)
        for sheet in writer.book.worksheets:
            sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
            from openpyxl.styles import Font,PatternFill
            for cell in sheet[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='214B75')
            for col in sheet.columns:
                cells=list(col)[:200]
                from openpyxl.utils import get_column_letter
                width=min(52,max(12,max(len(str(c.value or '')) for c in cells)+2))
                sheet.column_dimensions[get_column_letter(col[0].column)].width=width
    return buf.getvalue()
